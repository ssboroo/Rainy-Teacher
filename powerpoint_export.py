"""Create editable PowerPoint files for Rainy Teacher exports."""
from __future__ import annotations

import json
import re
import base64
from io import BytesIO
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


def _rgb(value: str) -> RGBColor:
    value = value.lstrip("#")
    return RGBColor(*(int(value[i:i + 2], 16) for i in (0, 2, 4)))


def _textbox(slide, text, x, y, w, h, *, size=18, color="#17243A", bold=False,
             align=PP_ALIGN.LEFT, font="Aptos", margin=0.02, valign=MSO_ANCHOR.MIDDLE):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(margin)
    frame.margin_top = frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.text = str(text or "")
    for run in paragraph.runs:
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = _rgb(color)
    return shape


def _rect(slide, x, y, w, h, fill=None, line=None, line_width=1, shape=MSO_SHAPE.RECTANGLE):
    item = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill:
        item.fill.solid()
        item.fill.fore_color.rgb = _rgb(fill)
    else:
        item.fill.background()
    if line:
        item.line.color.rgb = _rgb(line)
        item.line.width = Pt(line_width)
    else:
        item.line.fill.background()
    return item


def _line(slide, x1, y1, x2, y2, color, width=1):
    item = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    item.line.color.rgb = _rgb(color)
    item.line.width = Pt(width)
    return item


def _presentation(theme: dict, title: str, subject: str, slides: list[dict]) -> Presentation:
    accent = theme.get("accent", "#1C5D9D")
    bg = theme.get("bg", "#F5F8FC")
    ink = theme.get("ink", "#102B48")
    soft = theme.get("soft", "#DCEAF7")
    mode = theme.get("mode", "steppe")
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]

    cover = prs.slides.add_slide(blank)
    cover.background.fill.solid()
    cover.background.fill.fore_color.rgb = _rgb(bg)
    if mode in ("dark", "classic", "shape", "wave"):
        _rect(cover, 8.8, 0, 4.533, 7.5, soft)
        _rect(cover, 8.48, 0, .14, 7.5, accent)
        _rect(cover, 9.25, .75, 3.15, 3.15, None, accent, 1.2, MSO_SHAPE.OVAL)
        _rect(cover, 10.05, 1.55, 1.55, 1.55, None, accent, 1, MSO_SHAPE.OVAL)
    elif mode == "split":
        _rect(cover, 9.2, 0, 4.133, 7.5, accent)
        _rect(cover, 9.55, .45, 3.3, 6.6, soft)
        for i in range(4):
            _rect(cover, 10.0 + i * .52, 1.45 + i * .42, .26, .26, bg, accent, 1, MSO_SHAPE.DIAMOND)
    else:
        _rect(cover, 0, 0, .17, 7.5, accent)
        for i in range(5):
            _rect(cover, 10.6 + i * .4, .3, .18, .18, soft, accent, .7, MSO_SHAPE.DIAMOND)
    _textbox(cover, "RAINY TEACHER   ·   СУРГАЛТЫН ТАНИЛЦУУЛГА", .75, .55, 10.8, .32,
             size=10, color=accent, bold=True)
    _textbox(cover, subject or "БАГШИЙН PRESENTATION", .82, 2.05, 9.3, .42,
             size=13, color=accent, bold=True)
    title_size = 38 if len(title) < 42 else 31 if len(title) < 68 else 25
    _textbox(cover, title, .78, 2.58, 9.5, 2.0, size=title_size, color=ink, bold=True,
             valign=MSO_ANCHOR.MIDDLE)
    _rect(cover, .82, 4.87, 1.14, .065, accent)
    _textbox(cover, "Мэдлэгийг ойлгомжтой, сонирхолтой хүргэе.", .82, 5.16, 8.6, .52,
             size=17, color=ink)
    _line(cover, .82, 6.65, 12.45, 6.65, accent, .8)
    _textbox(cover, "БАГШИЙН ТУСЛАХ", .82, 6.76, 6, .27, size=9, color=ink, bold=True)
    _textbox(cover, "01", 11.8, 6.76, .62, .27, size=10, color=accent, bold=True, align=PP_ALIGN.RIGHT)

    expanded = []
    for slide in slides:
        stitle = str(slide.get("title", "Шинэ слайд")).strip() or "Шинэ слайд"
        elements = slide.get("elements")
        if isinstance(elements, list) and elements:
            expanded.append((stitle, [], elements, slide.get("background", bg)))
            continue
        points = slide.get("points", [])
        if isinstance(points, str):
            points = [part.strip() for part in points.splitlines() if part.strip()]
        points = [str(point).strip() for point in points if str(point).strip()] or ["Гол санаагаа энд бичнэ үү."]
        for part_no, start in enumerate(range(0, len(points), 4), 1):
            part_title = stitle if start == 0 else f"{stitle} · үргэлжлэл"
            expanded.append((part_title, points[start:start + 4], None, bg))

    for index, (stitle, points, elements, slide_bg) in enumerate(expanded, 2):
        page = prs.slides.add_slide(blank)
        page.background.fill.solid()
        page.background.fill.fore_color.rgb = _rgb(slide_bg)
        if elements is not None:
            for element in elements:
                try:
                    kind = element.get("type", "text")
                    # The canvas editor uses a 960×540 logical canvas.
                    x = float(element.get("x", 0)) * 13.333 / 960
                    y = float(element.get("y", 0)) * 7.5 / 540
                    w = float(element.get("w", 100)) * 13.333 / 960
                    h = float(element.get("h", 40)) * 7.5 / 540
                    if kind in ("rect", "ellipse"):
                        shape = MSO_SHAPE.OVAL if kind == "ellipse" else MSO_SHAPE.RECTANGLE
                        fill = element.get("fill", accent)
                        item = _rect(page, x, y, w, h, fill, None, shape=shape)
                    elif kind == "image" and element.get("data"):
                        raw = base64.b64decode(element["data"])
                        stream = _fit_image(raw)
                        scale = min(w / stream.width, h / stream.height)
                        iw, ih = stream.width * scale, stream.height * scale
                        page.shapes.add_picture(stream, Inches(x + (w-iw)/2), Inches(y + (h-ih)/2),
                                                width=Inches(iw), height=Inches(ih))
                    else:
                        size = max(8, min(96, int(float(element.get("font_size", 24)))))
                        _textbox(page, str(element.get("text", "")), x, y, w, h,
                                 size=size, color=element.get("color", ink),
                                 bold=bool(element.get("bold", False)), valign=MSO_ANCHOR.TOP)
                except (ValueError, TypeError, KeyError, OSError):
                    continue
            continue
        _rect(page, 0, 0, .12, 7.5, accent)
        _rect(page, 12.55, 0, .783, 7.5, soft)
        _rect(page, 12.77, .46, .26, .26, None, accent, 1.2, MSO_SHAPE.DIAMOND)
        _textbox(page, "RAINY TEACHER", .65, .42, 4.0, .3, size=10, color=accent, bold=True)
        _textbox(page, subject, 7.0, .42, 4.95, .3, size=10, color=accent, align=PP_ALIGN.RIGHT)
        _line(page, .65, .86, 11.95, .86, accent, .7)
        title_font = 29 if len(stitle) < 46 else 23 if len(stitle) < 70 else 19
        _textbox(page, stitle, .68, 1.13, 11.1, .83, size=title_font, color=ink, bold=True)
        _rect(page, .7, 2.02, .86, .055, accent)
        base_y = 2.38
        for n, point in enumerate(points):
            font_size = 16 if len(point) < 115 else 13
            height = .88
            y = base_y + n * 1.0
            _rect(page, .77, y + .14, .16, .16, accent, shape=MSO_SHAPE.OVAL)
            _textbox(page, point, 1.12, y, 10.75, height, size=font_size, color=ink, valign=MSO_ANCHOR.TOP)
        _line(page, .65, 6.83, 11.95, 6.83, accent, .65)
        _textbox(page, title, .65, 6.93, 9.5, .25, size=8, color=ink)
        _textbox(page, f"{index:02d}", 11.25, 6.9, .65, .3, size=10, color=accent, bold=True, align=PP_ALIGN.RIGHT)
    return prs


def save_presentation_pptx(deck: dict, path: str | Path, theme: dict) -> Path:
    slides = deck.get("slides", [])
    if isinstance(slides, str):
        try:
            slides = json.loads(slides)
        except (TypeError, json.JSONDecodeError):
            slides = []
    if not isinstance(slides, list):
        raise ValueError("Слайдын мэдээлэл буруу байна.")
    prs = _presentation(theme, str(deck.get("title", "Presentation")), str(deck.get("subject", "")), slides)
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(target))
    return target


def _fit_image(data: bytes) -> BytesIO:
    from PIL import Image
    image = Image.open(BytesIO(bytes(data)))
    image.load()
    original_width, original_height = image.size
    if image.mode not in ("RGB", "RGBA"):
        image = image.convert("RGBA")
    out = BytesIO()
    image.save(out, format="PNG", optimize=True)
    out.seek(0)
    out.width = original_width
    out.height = original_height
    return out


def save_certificate_pptx(record: dict, path: str | Path, designs: dict) -> Path:
    design = designs.get(record.get("style", "gold")) or designs.get("gold") or next(iter(designs.values()))
    accent, pale, paper = design["colors"]
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(11.69), Inches(8.27)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    try:
        custom = json.loads(record.get("design_json") or "{}")
    except (TypeError, json.JSONDecodeError):
        custom = {}
    if custom.get("elements"):
        custom_bg = custom.get("background", paper)
        slide.background.fill.fore_color.rgb = _rgb(custom_bg)
        values = {"{name}":str(record.get("recipient", "")),
                  "{achievement}":str(record.get("achievement", "")),
                  "{organization}":str(record.get("organization", "")),
                  "{teacher}":str(record.get("teacher", "")),
                  "{date}":str(record.get("issue_date", ""))}
        for element in custom["elements"]:
            try:
                kind = element.get("type", "text")
                x = float(element.get("x", 0)) * 11.69 / 960
                y = float(element.get("y", 0)) * 8.27 / 679
                w = float(element.get("w", 100)) * 11.69 / 960
                h = float(element.get("h", 40)) * 8.27 / 679
                if kind in ("rect", "ellipse"):
                    shape = MSO_SHAPE.OVAL if kind == "ellipse" else MSO_SHAPE.RECTANGLE
                    _rect(slide, x, y, w, h, element.get("fill", accent), shape=shape)
                elif kind == "image" and element.get("data"):
                    stream = _fit_image(base64.b64decode(element["data"]))
                    scale = min(w / stream.width, h / stream.height)
                    iw, ih = stream.width * scale, stream.height * scale
                    slide.shapes.add_picture(stream, Inches(x+(w-iw)/2), Inches(y+(h-ih)/2),
                                             width=Inches(iw), height=Inches(ih))
                else:
                    text = str(element.get("text", ""))
                    for token, value in values.items(): text = text.replace(token, value)
                    _textbox(slide, text, x, y, w, h,
                             size=max(8, min(96, int(float(element.get("font_size", 22))))),
                             color=element.get("color", "#1C2940"),
                             bold=bool(element.get("bold", False)), valign=MSO_ANCHOR.TOP)
            except (ValueError, TypeError, KeyError, OSError):
                continue
        target = Path(path)
        target.parent.mkdir(parents=True, exist_ok=True)
        prs.save(str(target))
        return target

    slide.background.fill.fore_color.rgb = _rgb(paper)

    # Double-line frame with stepped corner motifs inspired by Mongolian ulzii patterns.
    _rect(slide, .22, .22, 11.25, 7.83, None, accent, 2.1)
    _rect(slide, .34, .34, 11.01, 7.59, None, pale, .9)
    for sx in (1, -1):
        for sy in (1, -1):
            cx = .55 if sx == 1 else 11.14
            cy = .55 if sy == 1 else 7.72
            for offset, span in ((0, .72), (.16, .46), (.31, .24)):
                x_end = cx + sx * span
                y_end = cy + sy * span
                _line(slide, cx, cy + sy * offset, x_end, cy + sy * offset, accent, 1.3)
                _line(slide, cx + sx * offset, cy, cx + sx * offset, y_end, accent, 1.3)
            _rect(slide, cx - .10, cy - .10, .20, .20, paper, accent, 1.2, MSO_SHAPE.DIAMOND)
            _rect(slide, cx - .035, cy - .035, .07, .07, accent, shape=MSO_SHAPE.DIAMOND)

    if record.get("logo"):
        try:
            logo = _fit_image(record["logo"])
            scale = min(1.0 / logo.width, .62 / logo.height)
            logo_width, logo_height = logo.width * scale, logo.height * scale
            slide.shapes.add_picture(logo, Inches((11.69 - logo_width) / 2), Inches(.52),
                                     width=Inches(logo_width), height=Inches(logo_height))
            title_y = 1.23
        except Exception:
            title_y = .80
    else:
        title_y = .80

    _textbox(slide, str(record.get("title", "ӨРГӨМЖЛӨЛ")), .95, title_y, 9.8, .64,
             size=26, color=accent, bold=True, align=PP_ALIGN.CENTER)
    _textbox(slide, str(record.get("subtitle", "")), 1.2, title_y + .68, 9.3, .44,
             size=12, color="#627086", bold=True, align=PP_ALIGN.CENTER)
    _line(slide, 4.72, title_y + 1.22, 6.97, title_y + 1.22, pale, 1.4)

    recipient = str(record.get("recipient", ""))
    name_size = 31 if len(recipient) < 25 else 25 if len(recipient) < 38 else 20
    name_y = title_y + 1.54
    _textbox(slide, recipient, .95, name_y, 9.8, .75, size=name_size, color="#1C2940",
             bold=True, align=PP_ALIGN.CENTER)
    _line(slide, 2.65, name_y + .86, 9.05, name_y + .86, accent, 1.05)

    body_template = str(record.get("body", "{achievement}"))
    body = body_template.replace("{name}", recipient).replace("{achievement}", str(record.get("achievement", "")))
    _textbox(slide, body, 1.37, name_y + 1.12, 8.95, 1.20, size=14, color="#43516A",
             align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin=.06)

    footer_y = 6.61
    footer = [("СУРГУУЛЬ / БАЙГУУЛЛАГА", record.get("organization", "")),
              ("БАГШ", record.get("teacher", "")),
              ("ОЛГОСОН ОГНОО", record.get("issue_date", ""))]
    for i, (label, value) in enumerate(footer):
        x = .95 + i * 3.28
        _line(slide, x, footer_y, x + 2.45, footer_y, pale, 1)
        _textbox(slide, str(value or "________________"), x, footer_y + .08, 2.45, .28,
                 size=10, color="#43516A", align=PP_ALIGN.CENTER)
        _textbox(slide, label, x, footer_y + .39, 2.45, .25, size=8, color=accent,
                 bold=True, align=PP_ALIGN.CENTER)

    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(target))
    return target


def convert_pptx_to_legacy_ppt(pptx_path: str | Path, ppt_path: str | Path) -> Path:
    """Convert to the legacy binary .ppt format when Microsoft PowerPoint is installed."""
    try:
        import win32com.client
    except ImportError as exc:
        raise RuntimeError("Хуучин .ppt форматад хадгалахад Microsoft PowerPoint болон Office-ийн хөрвүүлэгч хэрэгтэй. .pptx форматыг сонгоно уу.") from exc
    app = presentation = None
    try:
        app = win32com.client.DispatchEx("PowerPoint.Application")
        app.Visible = 0
        presentation = app.Presentations.Open(str(Path(pptx_path).resolve()), ReadOnly=True, Untitled=False, WithWindow=False)
        presentation.SaveAs(str(Path(ppt_path).resolve()), 1)
    except Exception as exc:
        raise RuntimeError(".ppt болгон хөрвүүлж чадсангүй. Microsoft PowerPoint суусан эсэхийг шалгаад .pptx хэлбэрээр хадгална уу.") from exc
    finally:
        if presentation is not None:
            try:
                presentation.Close()
            except Exception:
                pass
        if app is not None:
            try:
                app.Quit()
            except Exception:
                pass
    return Path(ppt_path)
