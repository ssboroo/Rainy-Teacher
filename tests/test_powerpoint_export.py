import tempfile
import unittest
from pathlib import Path
from io import BytesIO

from PIL import Image
from pptx import Presentation
from powerpoint_export import save_certificate_pptx, save_presentation_pptx
from bagshiin_tuslah import presentation_html, certificate_html, Store


class PowerPointExportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.folder = Path(self.temp.name)

    def test_presentation_is_editable_and_preserves_every_point(self):
        deck = {
            "title": "Монгол хэлний хичээл",
            "subject": "Монгол хэл · 5-р анги",
            "slides": [{"title": "Эгшиг авиаг таних", "points": [
                "Эгшиг авиаг ялгана", "Үгийн үеийг тоолно", "Жишээгээр бататгана",
                "Хосоор ажиллана", "Дүгнэлт гаргана",
            ]}],
        }
        target = self.folder / "lesson.pptx"
        save_presentation_pptx(deck, target, {
            "accent": "#1C5D9D", "bg": "#F5F8FC", "ink": "#102B48",
            "soft": "#DCEAF7", "mode": "steppe",
        })
        reopened = Presentation(target)
        self.assertEqual(len(reopened.slides), 3)  # cover plus two editable content slides
        text = "\n".join(shape.text for slide in reopened.slides for shape in slide.shapes if shape.has_text_frame)
        for point in deck["slides"][0]["points"]:
            self.assertIn(point, text)
        self.assertEqual(reopened.slide_width, int(13.333 * 914400))

    def test_certificate_exports_editable_mongolian_text_and_frame(self):
        logo = BytesIO()
        Image.new("RGBA", (64, 64), (25, 60, 120, 255)).save(logo, format="PNG")
        record = {
            "style": "navy", "title": "ӨРГӨМЖЛӨЛ", "subtitle": "Амжилтаа үнэлүүлэв",
            "recipient": "Бат-Эрдэнэ", "achievement": "Улсын олимпиадад амжилттай оролцсон",
            "body": "{achievement} амжилт гаргасныг үнэлж {name}-д гардуулав.",
            "organization": "Эрдмийн оргил сургууль", "teacher": "Б. Саруул",
            "issue_date": "2026-10-15", "logo": logo.getvalue(), "logo_mime": "image/png",
        }
        target = self.folder / "certificate.pptx"
        save_certificate_pptx(record, target, {
            "navy": {"colors": ("#183F73", "#EFDFB3", "#FFFAF0")},
            "gold": {"colors": ("#A77920", "#F5E8BC", "#FFFAF0")},
        })
        reopened = Presentation(target)
        self.assertEqual(len(reopened.slides), 1)
        slide = reopened.slides[0]
        text = "\n".join(shape.text for shape in slide.shapes if shape.has_text_frame)
        self.assertIn("ӨРГӨМЖЛӨЛ", text)
        self.assertIn("Бат-Эрдэнэ", text)
        self.assertIn("Эрдмийн оргил сургууль", text)
        self.assertGreaterEqual(len(slide.shapes), 35)  # editable frame, corner ornaments, rules, and text
        pictures = [shape for shape in slide.shapes if shape.shape_type == 13]
        self.assertEqual(len(pictures), 1)
        self.assertAlmostEqual(pictures[0].width / pictures[0].height, 1.0, places=2)

    def test_custom_canvas_elements_export_as_individually_editable_shapes(self):
        image = BytesIO()
        Image.new("RGBA", (200, 100), (240, 170, 20, 255)).save(image, format="PNG")
        deck = {
            "title": "Canva маягийн засвар",
            "subject": "Байгалийн ухаан",
            "slides": [{"title": "Нарны аймаг", "points": [], "elements": [
                {"type": "text", "text": "НАРНЫ АЙМАГ", "x": 60, "y": 55, "w": 650, "h": 80,
                 "font_size": 32, "color": "#183F73", "bold": True},
                {"type": "rect", "x": 70, "y": 180, "w": 150, "h": 90, "fill": "#E4B946"},
                {"type": "image", "data": __import__("base64").b64encode(image.getvalue()).decode("ascii"),
                 "mime": "image/png", "x": 300, "y": 170, "w": 200, "h": 120},
            ]}],
        }
        target = self.folder / "custom-design.pptx"
        save_presentation_pptx(deck, target, {"accent":"#1C5D9D", "bg":"#FFFFFF", "ink":"#102B48", "soft":"#DCEAF7"})
        slide = Presentation(target).slides[1]
        texts = [shape.text for shape in slide.shapes if shape.has_text_frame]
        self.assertIn("НАРНЫ АЙМАГ", texts)
        self.assertEqual(sum(shape.shape_type == 13 for shape in slide.shapes), 1)
        self.assertEqual(len(slide.shapes), 3)
        html_preview = presentation_html({"title":"Нарны аймаг", "theme":"steppe", "slides":deck["slides"]})
        self.assertIn("НАРНЫ АЙМАГ", html_preview)
        self.assertIn("data:image/png;base64,", html_preview)
        self.assertIn("background:#E4B946", html_preview)

    def test_certificate_visual_layout_round_trips_to_pdf_and_powerpoint(self):
        record={"style":"navy","title":"БАТЛАМЖ","recipient":"Номин","achievement":"Улсын уралдаанд түрүүлсэн",
                "organization":"Эрдэм сургууль","teacher":"Б. Саруул","issue_date":"2026-10-15",
                "design_json":__import__("json").dumps({"background":"#FFFAF0","elements":[
                    {"type":"rect","x":20,"y":20,"w":920,"h":5,"fill":"#183F73"},
                    {"type":"text","text":"{name}","x":120,"y":210,"w":720,"h":70,"font_size":32,"color":"#183F73","bold":True},
                    {"type":"text","text":"{achievement}","x":120,"y":320,"w":720,"h":70,"font_size":18,"color":"#223344","bold":False},
                ]},ensure_ascii=False)}
        target=self.folder/"custom-certificate.pptx"
        save_certificate_pptx(record,target,{"navy":{"colors":("#183F73","#EFDFB3","#FFFAF0")}})
        slide=Presentation(target).slides[0]
        texts="\n".join(shape.text for shape in slide.shapes if shape.has_text_frame)
        self.assertIn("Номин",texts)
        self.assertIn("Улсын уралдаанд түрүүлсэн",texts)
        rendered=certificate_html(record)
        self.assertIn("Номин",rendered)
        self.assertIn("background:#FFFAF0",rendered)

    def test_custom_certificate_preserves_editorial_typography_in_html_and_pptx(self):
        import json
        record={"style":"navy","title":"ӨРГӨМЖЛӨЛ","recipient":"Номин","achievement":"Олимпиадын аварга",
                "organization":"Эрдэм сургууль","teacher":"Б. Саруул","issue_date":"2026-10-15",
                "design_json":json.dumps({"background":"#FFFAF0","elements":[
                    {"type":"text","text":"ӨРГӨМЖЛӨЛ","x":80,"y":70,"w":800,"h":70,"font":"Georgia","font_size":34,"color":"#183F73","bold":True,"align":"center"},
                    {"type":"text","text":"{name}","x":120,"y":210,"w":720,"h":70,"font":"Georgia","font_size":31,"color":"#26344B","bold":True,"align":"center"},
                ]},ensure_ascii=False)}
        html=certificate_html(record)
        self.assertIn("font-family:Georgia,serif",html)
        self.assertIn("text-align:center",html)
        self.assertIn("Хэвлэх / PDF хадгалах",html)
        target=self.folder/"certificate-type.pptx"
        save_certificate_pptx(record,target,{"navy":{"colors":("#183F73","#EFDFB3","#FFFAF0")}})
        slide=Presentation(target).slides[0]
        title=next(shape for shape in slide.shapes if shape.has_text_frame and shape.text=="ӨРГӨМЖЛӨЛ")
        self.assertEqual(title.text_frame.paragraphs[0].alignment, 2)  # PP_ALIGN.CENTER
        self.assertEqual(title.text_frame.paragraphs[0].runs[0].font.name,"Georgia")

    def test_database_adds_certificate_design_json_columns(self):
        store=Store(self.folder/"teacher.db")
        self.addCleanup(store.db.close)
        for table in ("certificate_templates","certificates"):
            columns={row[1] for row in store.db.execute(f"PRAGMA table_info({table})")}
            self.assertIn("design_json",columns)
        templates=store.all("SELECT design_json FROM certificate_templates")
        designs=[__import__("json").loads(row["design_json"]) for row in templates]
        self.assertEqual(len(designs),20)
        self.assertTrue(all(len(design["elements"])>=35 for design in designs))
        self.assertTrue(all(any(el.get("font")=="Georgia" and el.get("align")=="center" for el in design["elements"])
                            for design in designs))


if __name__ == "__main__":
    unittest.main()
